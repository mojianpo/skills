"""Extract bid-review rules from an xlsx workbook as JSON."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


HEADERS = {
    "审核模块": "audit_module",
    "审核项": "audit_item",
    "审核要点": "audit_point",
    "规则名称": "rule_name",
    "风险等级": "risk_level",
    "规则类型": "rule_type",
    "关联模版": "related_template",
    "规则描述": "rule_description",
}
REQUIRED_HEADERS = {"规则名称", "规则描述"}


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def parse_rules(path: Path, source_type: str) -> dict[str, Any]:
    if not path.is_file():
        raise FileNotFoundError(f"规则文件不存在: {path}")

    workbook = load_workbook(path, data_only=True, read_only=True)
    rules: list[dict[str, Any]] = []
    for worksheet in workbook.worksheets:
        rows = worksheet.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration:
            continue
        header_names = [_text(value) for value in header_row]
        if not REQUIRED_HEADERS.issubset(set(header_names)):
            missing = ", ".join(sorted(REQUIRED_HEADERS - set(header_names)))
            raise ValueError(f"工作表 {worksheet.title!r} 缺少必要表头: {missing}")
        indexes = {name: index for index, name in enumerate(header_names) if name in HEADERS}
        inherited = {field: "" for field in HEADERS.values()}
        for row_number, row in enumerate(rows, start=2):
            values = {}
            for name, field in HEADERS.items():
                index = indexes.get(name)
                values[field] = _text(row[index]) if index is not None and index < len(row) else ""
            for field, value in values.items():
                if value:
                    inherited[field] = value
            if not values["rule_name"] and not values["rule_description"]:
                continue
            rules.append({"sheet_name": worksheet.title, "row_number": row_number, **inherited})
    workbook.close()
    if not rules:
        raise ValueError("规则文件中没有有效规则")
    return {"rule_source": {"type": source_type, "path": str(path)}, "rules": rules}


def main() -> None:
    skill_root = Path(__file__).resolve().parents[1]
    default_path = skill_root / "data" / "rules.xlsx"
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--rules", type=Path, default=default_path)
    parser.add_argument("--source-type", choices=("user_file", "default_file"))
    args = parser.parse_args()
    source_type = args.source_type or ("default_file" if args.rules.resolve() == default_path.resolve() else "user_file")
    print(json.dumps(parse_rules(args.rules, source_type), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
