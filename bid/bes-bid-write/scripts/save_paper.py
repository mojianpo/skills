#!/usr/bin/env python3
"""Build a bid-response DOCX from constrained Markdown.

Supported Markdown: headings, paragraphs, bold text, unordered/ordered lists,
tables, local images, horizontal rules, and <!-- PAGE_BREAK -->.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any

try:
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt
except ImportError as exc:
    raise SystemExit(
        "缺少 python-docx。请先运行：python -m pip install -r requirements.txt"
    ) from exc


PLACEHOLDER_PATTERNS = (
    re.compile(r"【待补充[:：][^】]*】"),
    re.compile(r"\[(?:TODO|TBD|待补充)[^\]]*\]", re.IGNORECASE),
    re.compile(r"\b(?:TODO|TBD)\b", re.IGNORECASE),
)
IMAGE_RE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$")
HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
UNORDERED_RE = re.compile(r"^\s*[-*+]\s+(.+)$")
ORDERED_RE = re.compile(r"^\s*\d+[.)]\s+(.+)$")
BOLD_RE = re.compile(r"(\*\*.+?\*\*)")


def load_json(path: Path | None, default: dict[str, Any]) -> dict[str, Any]:
    if path is None:
        return default
    try:
        loaded = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ValueError(f"无法读取 JSON：{path}（{exc}）") from exc
    if not isinstance(loaded, dict):
        raise ValueError(f"JSON 顶层必须是对象：{path}")
    return loaded


def deep_merge(base: dict[str, Any], override: dict[str, Any]) -> dict[str, Any]:
    result = dict(base)
    for key, value in override.items():
        if isinstance(value, dict) and isinstance(result.get(key), dict):
            result[key] = deep_merge(result[key], value)
        else:
            result[key] = value
    return result


def default_config() -> dict[str, Any]:
    return {
        "page": {"top_mm": 25.4, "bottom_mm": 25.4, "left_mm": 30, "right_mm": 25.4},
        "fonts": {
            "body_east_asia": "宋体",
            "body_latin": "Times New Roman",
            "heading_east_asia": "黑体",
            "heading_latin": "Arial",
        },
        "sizes_pt": {
            "body": 12,
            "heading_1": 22,
            "heading_2": 16,
            "heading_3": 14,
            "heading_4": 12,
            "table": 10.5,
        },
        "paragraph": {"line_spacing": 1.5, "first_line_chars": 2, "space_after_pt": 0},
        "header": {"enabled": True},
        "footer": {"page_number": True},
    }


def set_run_font(run: Any, east_asia: str, latin: str, size: float, bold: bool | None = None) -> None:
    run.font.name = latin
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)


def configure_styles(document: Any, config: dict[str, Any]) -> None:
    fonts = config["fonts"]
    sizes = config["sizes_pt"]
    paragraph_cfg = config["paragraph"]

    normal = document.styles["Normal"]
    normal.font.name = fonts["body_latin"]
    normal.font.size = Pt(sizes["body"])
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), fonts["body_east_asia"])
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(paragraph_cfg["space_after_pt"])
    normal.paragraph_format.first_line_indent = Pt(
        sizes["body"] * paragraph_cfg["first_line_chars"]
    )

    for level in range(1, 5):
        style = document.styles[f"Heading {level}"]
        style.font.name = fonts["heading_latin"]
        style.font.size = Pt(sizes.get(f"heading_{level}", sizes["heading_4"]))
        style.font.bold = True
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), fonts["heading_east_asia"])
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.space_before = Pt(12 if level == 1 else 6)
        style.paragraph_format.space_after = Pt(6)

    if "Bid Table" not in [s.name for s in document.styles]:
        table_style = document.styles.add_style("Bid Table", WD_STYLE_TYPE.PARAGRAPH)
    else:
        table_style = document.styles["Bid Table"]
    table_style.font.name = fonts["body_latin"]
    table_style.font.size = Pt(sizes["table"])
    table_style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), fonts["body_east_asia"])
    table_style.paragraph_format.first_line_indent = Pt(0)
    table_style.paragraph_format.space_after = Pt(0)


def configure_section(section: Any, config: dict[str, Any]) -> None:
    page = config["page"]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(page["top_mm"])
    section.bottom_margin = Mm(page["bottom_mm"])
    section.left_margin = Mm(page["left_mm"])
    section.right_margin = Mm(page["right_mm"])
    section.different_first_page_header_footer = True


def add_field(paragraph: Any, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, text, end):
        run._r.append(element)


def add_page_number(paragraph: Any) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(paragraph, " PAGE ", "1")


def set_cell_shading(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def add_inline_markdown(paragraph: Any, text: str) -> None:
    for part in BOLD_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        else:
            paragraph.add_run(part)


def strip_front_matter(text: str) -> str:
    lines = text.splitlines()
    if lines and lines[0].strip() == "---":
        for index in range(1, len(lines)):
            if lines[index].strip() == "---":
                return "\n".join(lines[index + 1 :])
    return text


def split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_separator_row(line: str) -> bool:
    cells = split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def add_table(document: Any, rows: list[list[str]]) -> None:
    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for col_index in range(width):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            value = row[col_index] if col_index < len(row) else ""
            paragraph = cell.paragraphs[0]
            paragraph.style = "Bid Table"
            add_inline_markdown(paragraph, value)
            if row_index == 0:
                set_cell_shading(cell, "D9EAF7")
                for run in paragraph.runs:
                    run.bold = True
    document.add_paragraph()


def add_cover(document: Any, metadata: dict[str, Any], config: dict[str, Any]) -> None:
    fonts = config["fonts"]
    title = metadata.get("document_title", "投标响应文件")
    project = metadata.get("project_name", "")
    fields = [
        ("项目名称", project),
        ("招标编号", metadata.get("tender_number", "")),
        ("标段", metadata.get("lot", "")),
        ("投标人", metadata.get("bidder_name", "")),
        ("日期", metadata.get("date", "")),
    ]

    for _ in range(4):
        document.add_paragraph()
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run(project), fonts["heading_east_asia"], fonts["heading_latin"], 22, True)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    set_run_font(p.add_run(title), fonts["heading_east_asia"], fonts["heading_latin"], 30, True)
    for _ in range(5):
        document.add_paragraph()
    for label, value in fields[1:]:
        if not value:
            continue
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)
        set_run_font(
            p.add_run(f"{label}：{value}"),
            fonts["body_east_asia"],
            fonts["body_latin"],
            14,
        )
    confidentiality = metadata.get("confidentiality")
    if confidentiality:
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(24)
        set_run_font(
            p.add_run(str(confidentiality)),
            fonts["body_east_asia"],
            fonts["body_latin"],
            12,
            True,
        )
    document.add_page_break()


def add_toc(document: Any) -> None:
    heading = document.add_paragraph("目录", style="Heading 1")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Pt(0)
    add_field(paragraph, ' TOC \\o "1-3" \\h \\z \\u ', "请在 Word 中右键更新目录")
    document.add_page_break()


def add_header_footer(document: Any, metadata: dict[str, Any], config: dict[str, Any]) -> None:
    for section in document.sections:
        if config.get("header", {}).get("enabled", True):
            header = section.header.paragraphs[0]
            header.alignment = WD_ALIGN_PARAGRAPH.CENTER
            header.text = str(metadata.get("project_name", ""))
        if config.get("footer", {}).get("page_number", True):
            add_page_number(section.footer.paragraphs[0])


def render_markdown(document: Any, markdown: str, source_dir: Path) -> dict[str, int]:
    lines = strip_front_matter(markdown).splitlines()
    stats = {"headings": 0, "paragraphs": 0, "tables": 0, "images": 0}
    index = 0
    while index < len(lines):
        raw = lines[index]
        line = raw.strip()
        if not line:
            index += 1
            continue
        if line == "<!-- PAGE_BREAK -->":
            document.add_page_break()
            index += 1
            continue
        heading = HEADING_RE.match(line)
        if heading:
            level = min(len(heading.group(1)), 4)
            document.add_paragraph(heading.group(2), style=f"Heading {level}")
            stats["headings"] += 1
            index += 1
            continue
        image = IMAGE_RE.match(line)
        if image:
            image_path = (source_dir / image.group(2)).resolve()
            if not image_path.is_file():
                raise ValueError(f"图片不存在：{image.group(2)}")
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run().add_picture(str(image_path), width=Mm(145))
            if image.group(1):
                caption = document.add_paragraph(image.group(1))
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption.paragraph_format.first_line_indent = Pt(0)
            stats["images"] += 1
            index += 1
            continue
        if "|" in line and index + 1 < len(lines) and is_separator_row(lines[index + 1]):
            rows = [split_table_row(line)]
            index += 2
            while index < len(lines) and "|" in lines[index] and lines[index].strip():
                rows.append(split_table_row(lines[index]))
                index += 1
            add_table(document, rows)
            stats["tables"] += 1
            continue
        unordered = UNORDERED_RE.match(raw)
        ordered = ORDERED_RE.match(raw)
        if unordered or ordered:
            style = "List Bullet" if unordered else "List Number"
            paragraph = document.add_paragraph(style=style)
            paragraph.paragraph_format.first_line_indent = Pt(0)
            add_inline_markdown(paragraph, (unordered or ordered).group(1))
            stats["paragraphs"] += 1
            index += 1
            continue
        if line == "---":
            document.add_paragraph("─" * 40)
            index += 1
            continue

        paragraph_lines = [line]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if (
                not candidate
                or HEADING_RE.match(candidate)
                or IMAGE_RE.match(candidate)
                or UNORDERED_RE.match(lines[index])
                or ORDERED_RE.match(lines[index])
                or candidate == "<!-- PAGE_BREAK -->"
                or ("|" in candidate and index + 1 < len(lines) and is_separator_row(lines[index + 1]))
            ):
                break
            paragraph_lines.append(candidate)
            index += 1
        paragraph = document.add_paragraph()
        add_inline_markdown(paragraph, " ".join(paragraph_lines))
        stats["paragraphs"] += 1
    return stats


def find_placeholders(*texts: str) -> list[str]:
    found: list[str] = []
    for text in texts:
        for pattern in PLACEHOLDER_PATTERNS:
            found.extend(match.group(0) for match in pattern.finditer(text))
    return sorted(set(found))


def validate_docx(path: Path) -> list[str]:
    errors: list[str] = []
    if not path.is_file() or path.stat().st_size == 0:
        return ["DOCX 文件不存在或为空"]
    if not zipfile.is_zipfile(path):
        return ["输出不是有效的 OOXML/ZIP 文件"]
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        if "word/document.xml" not in names:
            errors.append("缺少 word/document.xml")
        if "[Content_Types].xml" not in names:
            errors.append("缺少 [Content_Types].xml")
    try:
        document = Document(str(path))
        body_text = "\n".join(p.text.strip() for p in document.paragraphs if p.text.strip())
        if not body_text:
            errors.append("DOCX 没有非空正文段落")
    except Exception as exc:  # python-docx exposes multiple parser errors
        errors.append(f"python-docx 无法重新打开输出：{exc}")
    return errors


def build(args: argparse.Namespace) -> dict[str, Any]:
    input_path = args.input.resolve()
    output_path = args.output.resolve()
    if not input_path.is_file():
        raise ValueError(f"输入文件不存在：{input_path}")
    markdown = input_path.read_text(encoding="utf-8-sig")
    if not markdown.strip():
        raise ValueError("输入 Markdown 为空")

    metadata = load_json(args.metadata.resolve() if args.metadata else None, {})
    config = deep_merge(
        default_config(),
        load_json(args.config.resolve() if args.config else None, {}),
    )
    metadata_text = json.dumps(metadata, ensure_ascii=False)
    placeholders = find_placeholders(markdown, metadata_text)
    if args.strict and placeholders:
        raise ValueError(f"严格模式发现 {len(placeholders)} 个待补占位符")

    document = Document()
    configure_styles(document, config)
    configure_section(document.sections[0], config)
    add_cover(document, metadata, config)
    add_toc(document)
    stats = render_markdown(document, markdown, input_path.parent)
    add_header_footer(document, metadata, config)

    settings = document.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    errors = validate_docx(output_path)
    report = {
        "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "input": str(input_path),
        "output": str(output_path),
        "output_bytes": output_path.stat().st_size if output_path.exists() else 0,
        "strict": args.strict,
        "placeholders": placeholders,
        "stats": stats,
        "validation_errors": errors,
        "status": "passed" if not errors and (not args.strict or not placeholders) else "failed",
        "notes": ["目录为 Word 域，首次打开文档后请更新全部域并进行视觉检查。"],
    }
    if args.report:
        args.report.resolve().parent.mkdir(parents=True, exist_ok=True)
        args.report.resolve().write_text(
            json.dumps(report, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
    if errors:
        raise ValueError("；".join(errors))
    return report


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="将标书 Markdown 生成并校验为 DOCX")
    parser.add_argument("--input", type=Path, required=True, help="UTF-8 Markdown 正文")
    parser.add_argument("--output", type=Path, required=True, help="输出 .docx 路径")
    parser.add_argument("--metadata", type=Path, help="封面字段 JSON")
    parser.add_argument("--config", type=Path, help="样式配置 JSON")
    parser.add_argument("--report", type=Path, help="JSON 校验报告")
    parser.add_argument("--strict", action="store_true", help="存在待补占位符时失败")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if args.output.suffix.lower() != ".docx":
        print("错误：--output 必须使用 .docx 扩展名", file=sys.stderr)
        return 2
    try:
        report = build(args)
    except (OSError, ValueError) as exc:
        print(f"错误：{exc}", file=sys.stderr)
        return 1
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
