#!/usr/bin/env python3
"""Validate the structure and unresolved placeholders of a DOCX file."""

from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from datetime import datetime
from pathlib import Path

try:
    from docx import Document
except ImportError as exc:
    raise SystemExit(
        "缺少 python-docx。请先运行：python -m pip install -r requirements.txt"
    ) from exc


PLACEHOLDER_RE = re.compile(
    r"【待补充[:：][^】]*】|\[(?:TODO|TBD|待补充)[^\]]*\]|\b(?:TODO|TBD)\b",
    re.IGNORECASE,
)


def inspect(path: Path) -> dict:
    errors: list[str] = []
    warnings: list[str] = []
    placeholders: list[str] = []
    stats = {"paragraphs": 0, "tables": 0, "images": 0}

    if not path.is_file() or path.stat().st_size == 0:
        errors.append("文件不存在或为空")
    elif not zipfile.is_zipfile(path):
        errors.append("不是有效的 OOXML/ZIP 文件")
    else:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            for required in ("[Content_Types].xml", "word/document.xml"):
                if required not in names:
                    errors.append(f"缺少 {required}")
            stats["images"] = sum(name.startswith("word/media/") for name in names)
        try:
            document = Document(str(path))
            paragraphs = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    paragraphs.extend(cell.text for cell in row.cells)
            stats["paragraphs"] = sum(bool(text.strip()) for text in paragraphs)
            stats["tables"] = len(document.tables)
            placeholders = sorted(set(PLACEHOLDER_RE.findall("\n".join(paragraphs))))
            if stats["paragraphs"] == 0:
                errors.append("没有非空正文")
            if placeholders:
                warnings.append(f"发现 {len(placeholders)} 个待补占位符")
        except Exception as exc:
            errors.append(f"无法解析 DOCX：{exc}")

    return {
        "checked_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "file": str(path.resolve()),
        "size_bytes": path.stat().st_size if path.exists() else 0,
        "status": "failed" if errors else ("warning" if warnings else "passed"),
        "errors": errors,
        "warnings": warnings,
        "placeholders": placeholders,
        "stats": stats,
        "notes": ["结构校验不能替代 Word 中的目录更新、分页和视觉检查。"],
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="校验标书 DOCX")
    parser.add_argument("docx", type=Path)
    parser.add_argument("--report", type=Path)
    parser.add_argument("--strict", action="store_true", help="占位符视为失败")
    args = parser.parse_args()
    report = inspect(args.docx)
    if args.strict and report["placeholders"]:
        report["errors"].append("严格模式不允许待补占位符")
        report["status"] = "failed"
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 1 if report["status"] == "failed" else 0


if __name__ == "__main__":
    raise SystemExit(main())
